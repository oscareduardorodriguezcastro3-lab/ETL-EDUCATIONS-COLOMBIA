"""Silver: contratos de datos, cuarentena y tablas con granularidad definida.

TXT/CSV se leen por bloques para no cargar 1,8 GB en memoria. Se conservan
sumas y conteos de Saber 11: promediar promedios daría resultados incorrectos.
"""

import logging
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .common import municipality, unique, write_csv


class Audit:
    """Registra balance por fuente y rechazos con ubicación recuperable en Bronze."""

    def __init__(self, run: Path):
        self.run = run
        self.counts = []

    def filter(self, frame, reasons, source, offset=0):
        """Una fila tiene un motivo principal; no se ocultan descartes."""
        bad = reasons.ne("")
        if bad.any():
            rejected = pd.DataFrame({
                "archivo": source, "registro": range(offset + 1, offset + len(frame) + 1),
                "motivo": reasons.to_numpy(),
            }).loc[bad.to_numpy()]
            path = self.run / "quality" / "rechazos.csv"
            path.parent.mkdir(parents=True, exist_ok=True)
            rejected.to_csv(path, index=False, mode="a", header=not path.exists())
        self.counts.append({"archivo": source, "leidos": len(frame),
                            "rechazados": int(bad.sum()), "validos": int((~bad).sum())})
        return frame.loc[~bad].copy()


def chunks(path, columns, size):
    """Valida cabecera y lee solo columnas necesarias, sin omitir errores de parseo."""
    header = pd.read_csv(path, sep=";", encoding="utf-8-sig", nrows=0).columns
    missing = set(columns) - set(header)
    if missing:
        raise ValueError(f"{path.name}: faltan columnas {sorted(missing)}")
    with pd.read_csv(path, sep=";", encoding="utf-8-sig", usecols=columns,
                     dtype="string", chunksize=size, keep_default_na=False,
                     on_bad_lines="error") as reader:
        yield from reader


def icfes(paths, years, size, audit):
    """Devuelve una fila por municipio-periodo con suma, conteo y media.

    Un identificador repetido dentro del periodo detiene la ejecución: no se
    elige arbitrariamente cuál puntaje conservar. Se revisa entre bloques.
    """
    parts, available = [], []
    columns = ["periodo", "estu_consecutivo", "cole_cod_mcpio_ubicacion", "punt_global"]
    for path in paths:
        match = re.fullmatch(r"Examen_Saber_11_(20\d{2}[12])\.txt", path.name)
        if not match:
            raise ValueError(f"Nombre ICFES inesperado: {path.name}")
        period = match[1]
        if int(period[:4]) not in years:
            continue
        if period in available:
            raise ValueError(f"Periodo repetido: {period}")
        available.append(period)
        logging.info("Silver ICFES | %s", period)
        offset, seen = 0, set()
        for raw in chunks(path, columns, size):
            raw = raw.reset_index(drop=True)
            frame = pd.DataFrame({"codigo_municipio": municipality(raw[columns[2]]),
                                  "puntaje": pd.to_numeric(raw.punt_global, errors="coerce")})
            reasons = pd.Series("", index=raw.index)
            reasons.loc[~frame.puntaje.between(0, 500)] = "puntaje_fuera_de_rango_o_nulo"
            reasons.loc[frame.codigo_municipio.isna()] = "codigo_invalido"
            reasons.loc[raw.periodo.str.strip().ne(period)] = "periodo_no_coincide_con_archivo"
            ids = raw.estu_consecutivo.str.strip()
            reasons.loc[ids.eq("")] = "identificador_vacio"
            present = ids[ids.ne("")]
            if present.duplicated().any() or any(value in seen for value in present):
                raise ValueError(f"{path.name}: identificador ICFES duplicado; revisar origen.")
            seen.update(present)
            valid = audit.filter(frame, reasons, path.name, offset)
            offset += len(raw)
            part = valid.groupby("codigo_municipio").puntaje.agg(
                suma_puntaje="sum", estudiantes="count").reset_index()
            part["periodo"] = period
            parts.append(part)
    if not parts:
        raise ValueError("No hay ICFES para los años configurados.")
    result = pd.concat(parts).groupby(["codigo_municipio", "periodo"], as_index=False)[
        ["suma_puntaje", "estudiantes"]].sum()
    result["anio"] = result.periodo.str[:4].astype(int)
    result["puntaje_promedio"] = result.suma_puntaje / result.estudiantes
    return result, sorted(available)


def crc(path, years, size, audit):
    """Suma segmentos residenciales por municipio, año y trimestre.

    Las filas fuera del alcance se cuentan aparte de las inválidas. No se
    deduplica por municipio: operadores, tecnologías y segmentos son aditivos.
    """
    parts, filtered = [], 0
    columns = ["ANNO", "TRIMESTRE", "ID_MUNICIPIO", "SEGMENTO", "ACCESOS"]
    offset = 0
    for raw in chunks(path, columns, size):
        raw = raw.reset_index(drop=True)
        frame = pd.DataFrame({"codigo_municipio": municipality(raw.ID_MUNICIPIO),
            "anio": pd.to_numeric(raw.ANNO, errors="coerce"),
            "trimestre": pd.to_numeric(raw.TRIMESTRE, errors="coerce"),
            "accesos": pd.to_numeric(raw.ACCESOS, errors="coerce")})
        reasons = pd.Series("", index=raw.index)
        reasons.loc[~frame.accesos.ge(0) | frame.accesos.mod(1).ne(0)] = "accesos_invalidos"
        reasons.loc[~frame.trimestre.isin([1, 2, 3, 4])] = "trimestre_invalido"
        reasons.loc[~frame.anio.between(1900, 2100) | frame.anio.mod(1).ne(0)] = "anio_invalido"
        reasons.loc[frame.codigo_municipio.isna()] = "codigo_invalido"
        valid = audit.filter(frame, reasons, path.name, offset)
        offset += len(raw)
        scope = valid.anio.isin(years) & raw.loc[valid.index, "SEGMENTO"].str.strip().str.lower().str.startswith("residencial")
        filtered += int((~scope).sum())
        valid = valid.loc[scope]
        parts.append(valid.groupby(["codigo_municipio", "anio", "trimestre"], as_index=False).accesos.sum())
    if not parts:
        raise ValueError("CSV CRC vacío.")
    result = pd.concat(parts).groupby(["codigo_municipio", "anio", "trimestre"], as_index=False).accesos.sum()
    result[["anio", "trimestre"]] = result[["anio", "trimestre"]].astype(int)
    return result, filtered


def dane(path, years, audit):
    """Lee Total municipal y transforma columnas de años en filas.

    Se valida la estructura real del anexo. No se suman Total, Cabecera y Rural:
    hacerlo contaría los mismos hogares dos veces.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Proyecciones Hogares mpio"]
        rows = list(sheet.iter_rows(values_only=True))
        if rows[8][2] != "Código Municipio" or rows[8][4] != "Área":
            raise ValueError("DANE: cambió la cabecera esperada en fila 9.")
        positions = {int(str(value)): i for i, value in enumerate(rows[9])
                     if str(value).isdigit()}
        if not set(years) <= positions.keys():
            raise ValueError("DANE no contiene todos los años solicitados.")
        data = []
        for row_number, row in enumerate(rows[11:], 12):
            if str(row[4]).strip() != "Total":
                continue
            for year in years:
                data.append({"codigo_municipio": str(row[2]).strip(),
                    "municipio": row[3], "departamento": row[1], "anio": year,
                    "hogares": row[positions[year]], "fila_excel": row_number})
    finally:
        workbook.close()
    frame = pd.DataFrame(data)
    if frame.empty:
        raise ValueError("DANE: no hay registros Total.")
    frame.codigo_municipio = municipality(frame.codigo_municipio)
    frame.hogares = pd.to_numeric(frame.hogares, errors="coerce")
    reasons = pd.Series("", index=frame.index)
    reasons.loc[~frame.hogares.gt(0) | frame.hogares.mod(1).ne(0)] = "hogares_invalidos"
    reasons.loc[frame.codigo_municipio.isna()] = "codigo_invalido"
    valid = audit.filter(frame, reasons, path.name)
    unique(valid, ["codigo_municipio", "anio"], "DANE")
    return valid


def build(files, config, run):
    """Ejecuta las tres transformaciones y persiste las tablas Silver."""
    audit = Audit(run)
    exam, periods = icfes(files["icfes"], config["years"], config["chunksize"], audit)
    logging.info("Silver CRC | accesos residenciales")
    internet, excluded = crc(files["crc"][0], config["years"], config["chunksize"], audit)
    households = dane(files["dane"][0], config["years"], audit)
    for name, table in [("icfes_municipio_periodo", exam), ("crc_municipio_trimestre", internet),
                        ("dane_municipio_anio", households)]:
        if table.empty:
            raise ValueError(f"Silver vacío: {name}")
        write_csv(table, run / "silver" / f"{name}.csv")
    balance = pd.DataFrame(audit.counts).groupby("archivo", as_index=False).sum()
    write_csv(balance, run / "quality" / "balance_fuentes.csv")
    return exam, internet, households, {"periodos_icfes": periods,
        "crc_validos_fuera_de_alcance": excluded,
        "rechazados": int(balance.rechazados.sum())}
